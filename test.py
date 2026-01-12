import json
from docx import Document
from docx.shared import Pt


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def safe_pt(value):
    try:
        return Pt(float(value)) if value is not None else None
    except Exception:
        return None


def merge_style_and_layout(styles, layout):
    merged = []

    for i, style_block in enumerate(styles):
        layout_block = layout[i] if i < len(layout) else {}

        merged.append({
            **style_block,
            "layout": layout_block
        })

    return merged


def build_docx(merged_data, output_path):
    doc = Document()

    for block in merged_data:
        p = doc.add_paragraph()

        # apply paragraph style
        try:
            p.style = block.get("style", "Normal")
        except Exception:
            p.style = "Normal"

        runs = block.get("runs", [])

        if not runs:
            p.add_run(block.get("text", ""))
            continue

        for r in runs:
            run = p.add_run(r.get("text", ""))

            if r.get("font"):
                run.font.name = r["font"]

            size = safe_pt(r.get("font_size_pt"))
            if size:
                run.font.size = size

            if r.get("bold") is not None:
                run.bold = r["bold"]

            if r.get("italic") is not None:
                run.italic = r["italic"]

    doc.save(output_path)
    print(f"DOCX created → {output_path}")


if __name__ == "__main__":
    styles = load_json("tmp/styles.json")
    layout = load_json("tmp/layout.json")

    merged = merge_style_and_layout(styles, layout)

    build_docx(
        merged_data=merged,
        output_path="tmp/reconstructed.docx"
    )
