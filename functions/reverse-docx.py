from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

def create_ufa_report_template():
    # Create a new document
    doc = Document()
    
    # Set up styles
    styles = doc.styles
    
    # Create Title style
    title_style = styles.add_style('CustomTitle', WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.name = 'Calibri'
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(12)
    
    # Create Subtitle style
    subtitle_style = styles.add_style('CustomSubtitle', WD_STYLE_TYPE.PARAGRAPH)
    subtitle_style.font.name = 'Calibri'
    subtitle_style.font.size = Pt(16)
    subtitle_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_style.paragraph_format.space_after = Pt(6)
    
    # Create Heading 1 style
    heading1_style = styles.add_style('CustomHeading1', WD_STYLE_TYPE.PARAGRAPH)
    heading1_style.font.name = 'Calibri'
    heading1_style.font.size = Pt(14)
    heading1_style.font.bold = True
    heading1_style.paragraph_format.space_before = Pt(18)
    heading1_style.paragraph_format.space_after = Pt(6)
    
    # Create Heading 2 style
    heading2_style = styles.add_style('CustomHeading2', WD_STYLE_TYPE.PARAGRAPH)
    heading2_style.font.name = 'Calibri'
    heading2_style.font.size = Pt(12)
    heading2_style.font.bold = True
    heading2_style.paragraph_format.space_before = Pt(12)
    heading2_style.paragraph_format.space_after = Pt(6)
    
    # Create Normal style
    normal_style = styles.add_style('CustomNormal', WD_STYLE_TYPE.PARAGRAPH)
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.space_after = Pt(6)
    
    # ==================== TITLE PAGE ====================
    
    # Add Title
    title = doc.add_paragraph()
    title.style = title_style
    title_run = title.add_run('Sample Report Template')
    title_run.font.size = Pt(36)  # Larger font for title
    title_run.font.bold = True
    
    # Add subtitle
    doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.style = subtitle_style
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run('Precision Micro-Boom Spray Bar Prototype for Ultra-Selective Ground Application')
    subtitle_run.font.size = Pt(18)
    
    # Add spacing
    for _ in range(3):
        doc.add_paragraph()
    
    # Add client information
    client_info = doc.add_paragraph()
    client_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    client_info.add_run('Client: United Farmers of Alberta (UFA)')
    client_run = client_info.runs[0]
    client_run.font.name = 'Calibri'
    client_run.font.size = Pt(14)
    
    client_info.add_run('\nPrepared for: Simon Cobban, United Farmers of Alberta')
    client_info.runs[1].font.name = 'Calibri'
    client_info.runs[1].font.size = Pt(14)
    
    client_info.add_run('\nPrepared by: Precision AI Engineering Team')
    client_info.runs[2].font.name = 'Calibri'
    client_info.runs[2].font.size = Pt(14)
    
    # Add date
    doc.add_paragraph()
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.add_run('Date: December 26, 2025')
    date_run.font.name = 'Calibri'
    date_run.font.size = Pt(14)
    
    # Add spacing
    for _ in range(3):
        doc.add_paragraph()
    
    # Add confidential notice
    confidential = doc.add_paragraph()
    confidential.alignment = WD_ALIGN_PARAGRAPH.CENTER
    confidential_run1 = confidential.add_run('Confidential and Proprietary')
    confidential_run1.italic = True
    confidential_run1.font.name = 'Calibri'
    confidential_run1.font.size = Pt(12)
    
    confidential.add_run('\n')
    confidential_run2 = confidential.add_run('For discussion purposes only')
    confidential_run2.italic = True
    confidential_run2.font.name = 'Calibri'
    confidential_run2.font.size = Pt(12)
    
    # Add page break
    doc.add_page_break()
    
    # ==================== TABLE OF CONTENTS ====================
    
    # Add Table of Contents heading
    toc_heading = doc.add_paragraph('Table of Contents')
    toc_heading.style = heading1_style
    
    # Add spacing
    doc.add_paragraph()
    
    # Add TOC entries (simulated since python-docx doesn't have built-in TOC)
    toc_items = [
        ('Table of Contents', '2'),
        ('Executive Summary', '3'),
        ('HEADING STYLE GOES HERE', '3'),
        ('Objectives', '3')
    ]
    
    for item_text, page_num in toc_items:
        toc_item = doc.add_paragraph()
        
        # Create the text part
        text_run = toc_item.add_run(item_text)
        text_run.font.name = 'Calibri'
        text_run.font.size = Pt(11)
        
        # Add tab and dots
        toc_item.add_run('\t')
        
        # Calculate dots length (this is a simple simulation)
        dots_needed = 80 - len(item_text)
        if dots_needed > 0:
            dots_run = toc_item.add_run('.' * dots_needed)
            dots_run.font.name = 'Calibri'
            dots_run.font.size = Pt(11)
        
        # Add page number
        page_run = toc_item.add_run(page_num)
        page_run.font.name = 'Calibri'
        page_run.font.size = Pt(11)
    
    # Add page break
    doc.add_page_break()
    
    # ==================== EXECUTIVE SUMMARY ====================
    
    # Add Executive Summary
    exec_summary_heading = doc.add_paragraph('Executive Summary')
    exec_summary_heading.style = heading1_style
    
    # Add Executive Summary content
    exec_content = doc.add_paragraph()
    exec_content.style = normal_style
    exec_run1 = exec_content.add_run('Lorem Ipsum')
    exec_run1.bold = True
    exec_content.add_run(' is simply dummy text of the printing and typesetting industry. Lorem Ipsum has been the industry\'s standard dummy text ever since the 1500s, when an unknown printer took a galley of type and scrambled it to make a type specimen book. It has survived not only five centuries, but also the leap into electronic typesetting, remaining essentially unchanged. It was ')
    exec_content.add_run('popularised')
    exec_content.add_run(' in the 1960s with the release of Letraset sheets containing Lorem Ipsum passages, and more recently with desktop publishing software like Aldus PageMaker including versions of Lorem Ipsum.')
    
    # Add spacing
    doc.add_paragraph()
    
    # ==================== HEADING STYLE SECTION ====================
    
    # Add Heading Style Goes Here
    heading_style = doc.add_paragraph('HEADING STYLE GOES HERE')
    heading_style.style = heading1_style
    
    # Add content
    heading_content = doc.add_paragraph()
    heading_content.style = normal_style
    heading_content.add_run('Lorem Ipsum')
    heading_content.add_run(' is simply dummy text of the printing and typesetting industry. Lorem Ipsum has been the industry\'s standard dummy text ever since the 1500s, when an unknown printer took a galley of type and scrambled it to make a type specimen book. It has survived not only five centuries, but also the leap into electronic typesetting, remaining essentially unchanged. It was ')
    heading_content.add_run('popularised')
    heading_content.add_run(' in the 1960s with the release of Letraset sheets containing Lorem Ipsum passages, and more recently with desktop publishing software like Aldus PageMaker including versions of Lorem Ipsum.')
    
    # Add spacing
    doc.add_paragraph()
    
    # ==================== OBJECTIVES SECTION ====================
    
    # Add Objectives heading
    objectives_heading = doc.add_paragraph('Objectives')
    objectives_heading.style = heading2_style
    
    # Add bullet points
    bullet_points = [
        'You can put a few bullet points here',
        'You can put a few bullet points here', 
        'You can put a few bullet points here',
        'You can put a few bullet points here',
        'You can put a few bullet points here'
    ]
    
    for point in bullet_points:
        para = doc.add_paragraph(style='List Bullet')
        run = para.add_run(point)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
    
    # ==================== FOOTER ====================
    
    # Add footer with company information
    section = doc.sections[0]
    footer = section.footer
    
    # Clear default footer paragraph
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    
    # Add company info to footer
    footer_run = footer_paragraph.add_run('1965 Broad Street, Regina, SK, S4P 1Y1\n306.581.6888\ninfo@precision.ai')
    footer_run.font.name = 'Calibri'
    footer_run.font.size = Pt(9)
    
    # ==================== SAVE DOCUMENT ====================
    
    # Save the document
    filename = "Precision_AI_UFA_Report_Template.docx"
    doc.save(filename)
    
    print(f"✅ Document created successfully: {filename}")
    print(f"📁 File saved in: {os.path.abspath(filename)}")
    print("\nDocument includes:")
    print("- Title page with client information")
    print("- Table of Contents")
    print("- Executive Summary")
    print("- Heading style section")
    print("- Objectives with bullet points")
    print("- Footer with company details")

# Run the function
if __name__ == "__main__":
    create_ufa_report_template()